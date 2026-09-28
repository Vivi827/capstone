# -*- coding: utf-8 -*-
"""Export the rubric v2.2 scoring workbook (4 reviewer slots).

Schema follows docs/AI_TACTIC/Pally_RUBRIC_v2.2_md (PR #82) exactly:
speech function (10 codes) -> language_form -> Formality/Humor/Curiosity
score + case_type + score_status + confidence + evidence per axis.
Energy/Intimacy stay null (no audio, no usage history in this experiment).

Item source: data/fixtures/humor_v2.2_candidates_check.xlsx (remainder50
text reused + hand-picked boundary-candidate additions). Old remainder50
scores are NOT carried over -- v2.2 forbids reusing v2.1-era labels as-is.

Run from the repository root:
  python scripts/export_humor_v2_2_workbook.py
"""

from __future__ import annotations

import random
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

CANDIDATES_PATH = Path("data/fixtures/humor_v2.2_candidates_check.xlsx")
OUT_DIR = Path("data/fixtures")
SLOTS = ("A", "B", "C", "D")
SEED = 20260921

HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
GROUP_FILL = PatternFill("solid", fgColor="EFEFEF")
BOLD = Font(bold=True)

FUNCTIONS = [
    ("inform_describe", "사실 전달·상황 묘사", "The train leaves at six."),
    ("self_disclose", "개인 경험·상태 공유", "I felt nervous during the interview."),
    ("opinion_evaluate", "의견·평가", "That ending was surprisingly good."),
    ("ask_fact_meaning", "사실·정의 질문", "What does maintain mean?"),
    ("ask_explain_explore", "이유·설명·탐색 질문", "Why does that happen?"),
    ("confirm_clarify", "확인·명료화", "Do you mean the first one?"),
    ("request_direct", "행동 요청·지시·제안", "Could you open the file?"),
    ("social_phatic", "인사·안부·관계 유지", "How are you? / How about you?"),
    ("answer_acknowledge", "답변·동의·백채널", "Yes, that makes sense. / Really?"),
    ("express_react", "감정·반응 표현", "That's amazing!"),
]
FUNCTION_CODES = [f[0] for f in FUNCTIONS]

LANGUAGE_FORMS = [
    ("standard_or_unmarked", "뚜렷한 오류·비표준성이 없음"),
    ("likely_learner_error", "의미는 전달되나 비의도적 학습자 오류로 보임"),
    ("conventional_nonstandard", "ain't, 축약, 구어 문법처럼 관습화된 비표준형"),
    ("playful_deviation", "의도적 철자·문법 파괴, 밈 문형, 과장된 반복 등 유희적 변형의 근거가 있음"),
    ("intent_unclear", "오류인지 의도적 변형인지 단일 발화로 판정 불가"),
]
LANGUAGE_FORM_CODES = [f[0] for f in LANGUAGE_FORMS]

CASE_TYPES = ["clear", "boundary", "context_insufficient"]
SCORE_STATUS = ["final", "provisional"]

AXES = ("Formality", "Humor", "Curiosity")

RUBRIC_ANCHORS = {
    "Formality": [
        (0, "“Fix this for me.”", "직접 명령이며 완화와 상대 고려가 거의 없음"),
        (25, "“Can you fix this for me?”", "요청 형식이나 완화가 제한적"),
        (50, "“Could you help me fix this?”", "중립적인 회화의 정중한 요청"),
        (75, "“Could you help me fix this when you have a moment?”", "상대의 시간을 고려하고 요구를 완화"),
        (100, "“Sorry to bother you, but would you mind helping me fix this when it's convenient?”", "사과·허락·상대 편의를 함께 명시"),
    ],
    "Humor": [
        (0, "“I stayed home and watched a movie.”", "유희·비문자성 단서가 없는 문자적 진술"),
        (25, "“Well, that went great.”", "반어 가능성이 있으나 단일 문장에서는 약하고 모호함"),
        (50, "“My alarm clock and I are in a toxic relationship.”", "명확한 의인화와 가벼운 불일치가 있음"),
        (75, "“Girl, you ate—and left absolutely no crumbs.”", "검수 가능한 슬랭 의미와 과장된 유희적 칭찬이 결합됨"),
        (100, "“My deadline saw me relaxing and said, 'Absolutely not,' then spawned three side quests.”", "의인화·반전·과장·확장된 유희 프레이밍이 강하게 결합됨"),
    ],
    "Curiosity": [
        (0, "“I went hiking yesterday.”", "질문·정보 요청·탐색 표지가 없음"),
        (25, "“Can you open the file?”", "의문문 형태지만 주기능은 행동 요청"),
        (50, "“Is this the correct answer?”", "명시적인 확인 질문이며 답의 범위가 제한됨"),
        (75, "“What do you think would work better?”", "관점과 대안을 요구하고 설명을 유도"),
        (100, "“Why does this happen, and how would it change if we removed this step?”", "이유·메커니즘·가정적 후속 탐색이 한 발화에 결합됨"),
    ],
}

COMPONENT_WEIGHTS = {
    "Formality": [
        ("상대 존중·체면 고려", "40%", "비하·압박·상대 무시", "중립적이거나 대인 행위 없음", "상대 입장·편의·체면을 명시적으로 고려"),
        ("직접성 완화", "35%", "직접 명령·강한 요구", "중립적 표현 또는 통상적 요청", "허락·선택권·부담 완화를 강하게 제공"),
        ("구어 register", "25%", "매우 캐주얼·비격식", "일상적 중립 회화", "매우 격식적·공식적 구어"),
    ],
    "Humor": [
        ("유희적 의도·프레이밍", "40%", "문자적·중립적 전달", "장난 가능성이 있으나 불명확", "유희적 전달 의도가 강하고 일관됨"),
        ("비문자성·불일치·반전", "30%", "문자 의미 그대로", "약한 과장·중의성", "말장난·아이러니·강한 반전·예상 위반"),
        ("관습적 유머·밈·슬랭 단서", "20%", "관련 단서 없음", "표현은 있으나 의미·용도가 불확실", "검수된 표현을 유희적 의미로 명확히 사용"),
        ("표현 강화", "10%", "강화 방식 없음", "일부 반복·리듬·표기 변형", "유희를 분명히 강화하는 의도적 형식 사용"),
    ],
    "Curiosity": [
        ("탐색 요구 강도", "50%", "답이나 정보 요구 없음", "제한된 확인·반응 요구", "정보·설명·관점을 명시적으로 요구"),
        ("탐색 깊이", "25%", "탐색 없음", "단일 사실·의미 확인", "이유·원리·비교·가정까지 탐색"),
        ("답변 개방성", "15%", "실질적 답이 필요 없음", "짧고 제한된 답이면 충분", "설명·근거·복수 가능성을 요구"),
        ("인식적 표지", "10%", "모름·궁금함 표지 없음", "정보 부족이 간접적으로 드러남", "알고 싶은 점·불확실성을 명시"),
    ],
}

CURIOSITY_CLAMP = [
    ("비질문 진술", "0-15"),
    ("수사 질문·단순 되묻기", "10-30"),
    ("행동 요청", "10-30"),
    ("확인 질문", "25-50"),
    ("의례적 안부·대화권 반환", "30-55"),
    ("사실·단어 의미 질문", "45-70"),
    ("의견·선택 요청", "55-80"),
    ("이유·원리·과정 질문", "70-95"),
    ("한 발화 안의 가정·비교·후속 탐색", "85-100"),
]


def load_items() -> list[dict]:
    wb = openpyxl.load_workbook(CANDIDATES_PATH, data_only=True)
    ws = wb["후보"]
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        items.append({"item_id": row[0], "source": row[1], "utterance": row[2]})
    return items


def sheet_guide(wb: openpyxl.Workbook, slot: str, n_items: int, seed: int) -> None:
    ws = wb.create_sheet("안내")
    lines = [
        "Pally rubric v2.2 채점 안내 (단일 발화, 3축)",
        "",
        f"검수자 슬롯: {slot}   /   문항 수: {n_items}   /   순서 seed: {seed}",
        "",
        "[ 목적 ]",
        "이 실험의 주 평가축은 Humor다. Formality/Curiosity는 보조 결과로 같이 수집한다.",
        "Energy와 Intimacy는 이번 실험에서 측정하지 않는다 (실제 음성·사용 이력이 없음).",
        "",
        "[ 채점 순서 ]",
        "1. '발화기능·문법기준' 시트에서 10개 발화 기능과 5개 language_form을 먼저 읽는다.",
        "2. '채점기준표'와 '채점 예시' 시트로 Formality/Humor/Curiosity 판단 기준을 확인한다.",
        "3. '채점 시트'에서 각 발화에 대해:",
        "   a. 주 발화 기능 1개 선택 (필요하면 보조 기능 1개 추가)",
        "   b. language_form 선택",
        "   c. Formality/Humor/Curiosity 점수(0~100)와 case_type, score_status, confidence(0~4), 근거를 기록",
        "4. 상단 reviewer_id 칸에 본인 식별자를 입력한다.",
        "",
        "[ 규칙 (v2.2 원문 그대로) ]",
        "- 직전·다음 턴, 출처 corpus, 모델 점수, 다른 평가자의 점수를 보지 않는다. 현재 발화 하나만 본다.",
        "- 0은 축의 낮은 끝을 관찰했다는 뜻이고, null은 해당 입력으로 측정하지 않았다는 뜻이다 (Energy/Intimacy만 null).",
        "- 문법 오류 자체는 세 축 모두 직접 점수에 반영하지 않는다 (language_form으로 별도 기록).",
        "- joke/humorous를 발화 기능으로 쓰지 않는다 (Humor 정답이 기능 선택에 순환 반영되는 것을 막기 위함).",
        "- 여러 해석이 가능하면 case_type을 boundary 또는 context_insufficient로 표시하고, 그래도 가장 지지되는 잠정 점수를 기록한다 (판단 불가 시 50점 기본값 금지).",
        "- 다른 검수자와 상의하지 말고 독립적으로 채점한다.",
    ]
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 100
    for cell in ws["A"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=13)


def sheet_functions(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("발화기능·문법기준")
    ws.append(["공통 발화 기능 10개"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["코드", "발화 기능", "예시"])
    for cell in ws[2]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    for code, label, example in FUNCTIONS:
        ws.append([code, label, example])
    ws.append([])
    row0 = ws.max_row + 1
    ws.append(["문법 오류·비표준 표현 (language_form)"])
    ws[f"A{row0}"].font = Font(bold=True, size=12)
    ws.append(["코드", "정의"])
    for cell in ws[row0 + 1]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    for code, definition in LANGUAGE_FORMS:
        ws.append([code, definition])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 46
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def sheet_rubric(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("채점기준표")
    r = 1
    for axis in AXES:
        ws.cell(row=r, column=1, value=f"{axis} — 구성요소 가중치").font = Font(bold=True, size=12)
        r += 1
        ws.append(["구성요소", "가중치", "수준 0", "수준 2", "수준 4"])
        for cell in ws[r]:
            cell.font = BOLD
            cell.fill = HEADER_FILL
        r += 1
        for comp in COMPONENT_WEIGHTS[axis]:
            ws.append(list(comp))
            r += 1
        r += 1
    ws.cell(row=r, column=1, value="Curiosity 주 발화 기능별 허용 점수 범위 (clamp)").font = Font(bold=True, size=12)
    r += 1
    ws.append(["주 발화 기능", "허용 점수 범위"])
    for cell in ws[r + 1]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    r += 1
    for func, rng in CURIOSITY_CLAMP:
        ws.append([func, rng])
    r = ws.max_row + 2
    ws.cell(row=r, column=1, value="경계 사례·점수 상태").font = Font(bold=True, size=12)
    r += 1
    ws.append(["case_type", "의미"])
    for cell in ws[r + 1]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    ws.append(["clear", "현재 발화에서 우세한 의미와 축 점수를 안정적으로 판단 가능"])
    ws.append(["boundary", "복수 해석이 가능하지만 현재 발화 안에서 우세한 판단은 가능"])
    ws.append(["context_insufficient", "문맥 없이는 우세한 의미나 의도를 결정하기 어려움"])
    ws.append([])
    ws.append(["score_status", "의미"])
    for cell in ws[ws.max_row]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    ws.append(["final", "현재 발화에서 점수를 확정할 근거가 충분함"])
    ws.append(["provisional", "숫자는 기록하지만 문맥 부족으로 잠정적임"])

    for col, w in zip("ABCDE", [26, 34, 30, 30, 30]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def sheet_examples(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("채점 예시")
    r = 1
    for axis in AXES:
        ws.cell(row=r, column=1, value=f"{axis} anchor 예시").font = Font(bold=True, size=12)
        r += 1
        ws.append(["점수", "예시 발화", "판단 근거"])
        for cell in ws[r]:
            cell.font = BOLD
            cell.fill = HEADER_FILL
        r += 1
        for score, utt, reason in RUBRIC_ANCHORS[axis]:
            ws.append([score, utt, reason])
            r += 1
        r += 1

    ws.cell(row=r, column=1, value="완성된 채점 레코드 예시 (v2.2 스키마)").font = Font(bold=True, size=12)
    r += 1
    example_lines = [
        'target_utterance: "pls halp, my brain has left the chat"',
        "primary_function: request_direct   /   secondary_function: express_react",
        "language_form: playful_deviation",
        "Formality=25 (case_type=clear, status=final)",
        "Humor=80 (case_type=clear, status=final) — 근거: 인터넷 축약과 캐주얼 register, 의도적 표기 변형",
        "Curiosity=15 (case_type=clear, status=final) — 근거: 정보 탐색보다 도움 요청이 주기능",
        "Energy=null, Intimacy=null",
    ]
    for line in example_lines:
        ws.cell(row=r, column=1, value=line)
        r += 1

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 60
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


AXIS_COL_GROUPS = {
    "Formality": ["Formality", "Formality_case_type", "Formality_status", "Formality_confidence", "Formality_evidence"],
    "Humor": ["Humor", "Humor_case_type", "Humor_status", "Humor_confidence", "Humor_evidence"],
    "Curiosity": ["Curiosity", "Curiosity_case_type", "Curiosity_status", "Curiosity_confidence", "Curiosity_evidence"],
}

BASE_COLS = ["item_id", "utterance", "primary_function", "secondary_function", "function_ambiguous", "language_form"]
TAIL_COLS = ["Energy", "Intimacy", "annotator_id", "메모(선택)"]


def sheet_scoring(wb: openpyxl.Workbook, slot: str, items: list[dict], seed: int) -> None:
    ws = wb.create_sheet("채점 시트")
    ws.append([f"reviewer_id →", "", "", "", "(여기에 본인 식별자 입력)"])
    ws.append([f"reviewer_slot →", slot])
    ws.append([])

    columns = BASE_COLS + AXIS_COL_GROUPS["Formality"] + AXIS_COL_GROUPS["Humor"] + AXIS_COL_GROUPS["Curiosity"] + TAIL_COLS
    header_row = ws.max_row + 1
    ws.append(columns)
    for cell in ws[header_row]:
        cell.font = BOLD
        cell.fill = HEADER_FILL

    order = list(range(len(items)))
    random.Random(f"{seed}|humor_v2.2|{slot}").shuffle(order)

    col_index = {name: i + 1 for i, name in enumerate(columns)}

    for idx in order:
        item = items[idx]
        row = [""] * len(columns)
        row[col_index["item_id"] - 1] = item["item_id"]
        row[col_index["utterance"] - 1] = item["utterance"]
        row[col_index["Energy"] - 1] = "null"
        row[col_index["Intimacy"] - 1] = "null"
        ws.append(row)

    first_data_row = header_row + 1
    last_data_row = header_row + len(items)

    def col_letter(name: str) -> str:
        idx = col_index[name]
        return openpyxl.utils.get_column_letter(idx)

    def col_range(name: str) -> str:
        letter = col_letter(name)
        return f"{letter}{first_data_row}:{letter}{last_data_row}"

    dv_function = DataValidation(type="list", formula1=f'"{",".join(FUNCTION_CODES)}"', allow_blank=True)
    dv_function_secondary = DataValidation(type="list", formula1=f'"{",".join(FUNCTION_CODES)}"', allow_blank=True)
    dv_bool = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    dv_langform = DataValidation(type="list", formula1=f'"{",".join(LANGUAGE_FORM_CODES)}"', allow_blank=True)
    dv_case_type = DataValidation(type="list", formula1=f'"{",".join(CASE_TYPES)}"', allow_blank=True)
    dv_status = DataValidation(type="list", formula1=f'"{",".join(SCORE_STATUS)}"', allow_blank=True)
    dv_confidence = DataValidation(type="whole", operator="between", formula1=0, formula2=4, allow_blank=True)
    dv_score = DataValidation(type="whole", operator="between", formula1=0, formula2=100, allow_blank=True)

    dv_function.add(col_range("primary_function"))
    dv_function_secondary.add(col_range("secondary_function"))
    dv_bool.add(col_range("function_ambiguous"))
    dv_langform.add(col_range("language_form"))

    for axis in AXES:
        dv_score.add(col_range(axis))
        dv_case_type.add(col_range(axis + "_case_type"))
        dv_status.add(col_range(axis + "_status"))
        dv_confidence.add(col_range(axis + "_confidence"))

    for dv in (dv_function, dv_function_secondary, dv_bool, dv_langform, dv_case_type, dv_status, dv_confidence, dv_score):
        ws.add_data_validation(dv)

    widths = {
        "item_id": 16, "utterance": 55, "primary_function": 20, "secondary_function": 20,
        "function_ambiguous": 10, "language_form": 22,
    }
    for axis in AXES:
        widths.update({
            axis: 8, f"{axis}_case_type": 20, f"{axis}_status": 14,
            f"{axis}_confidence": 12, f"{axis}_evidence": 34,
        })
    widths.update({"Energy": 8, "Intimacy": 8, "annotator_id": 14, "메모(선택)": 30})
    for name, w in widths.items():
        ws.column_dimensions[col_letter(name)].width = w

    for row in ws.iter_rows(min_row=first_data_row, max_row=last_data_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_workbook(slot: str, items: list[dict], seed: int, out_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet_guide(wb, slot, len(items), seed)
    sheet_functions(wb)
    sheet_rubric(wb)
    sheet_examples(wb)
    sheet_scoring(wb, slot, items, seed)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> None:
    items = load_items()
    print(f"loaded {len(items)} candidate items")
    for slot in SLOTS:
        out_path = OUT_DIR / f"humor_v2.2_scoring_slot{slot}.xlsx"
        build_workbook(slot, items, SEED, out_path)
        print(f"  wrote {out_path}")


if __name__ == "__main__":
    main()
