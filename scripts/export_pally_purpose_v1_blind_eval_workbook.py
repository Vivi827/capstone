# -*- coding: utf-8 -*-
"""Export the Phase 3 blind A/B evaluation workbook (4 researcher slots),
per PM final design Sec. 12 (blind procedure) and Sec. 13 (scoring rubric).

Each of the 20 final-test items' A/B responses are anonymized as X/Y with a
per-evaluator, per-item randomized order (seeded, reproducible). Condition
labels, predicted purpose, axis scores, and reference_response are NEVER
shown to evaluators (Sec. 12.1). good_response and winner are computed by
formula from the rubric (Sec. 13.3/13.4), not typed in by hand, so the
mechanical rule is applied consistently across all 4 raters.

The X/Y -> condition mapping is written to a separate internal file
(NOT distributed to evaluators) for later de-anonymization when scoring.

Usage:
  python scripts/export_pally_purpose_v1_blind_eval_workbook.py
"""
from __future__ import annotations

import json
import random
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
ITEMS_PATH = ROOT / "data/fixtures/pally_purpose_v1_items_100.jsonl"
RESPONSES_PATH = ROOT / "data/fixtures/pally_purpose_v1_responses.jsonl"
OUT_DIR = ROOT / "data/fixtures"
MAPPING_PATH = ROOT / "data/fixtures/pally_purpose_v1_blind_eval_mapping.json"  # internal, gitignored
SLOTS = ("A", "B", "C", "D")
SEED = 20260929

HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
BOLD = Font(bold=True)

DIMENSIONS = [
    ("의미와_태도_보존", 4),
    ("목적_충족", 3),
    ("스타일_적절성", 3),
    ("대화_연속성", 3),
    ("관련성과_자연스러움", 3),
]
YN = ["Y", "N"]


def load_items():
    return {it["item_id"]: it for it in (json.loads(l) for l in ITEMS_PATH.read_text(encoding="utf-8").splitlines() if l.strip())}


def load_responses():
    rows = [json.loads(l) for l in RESPONSES_PATH.read_text(encoding="utf-8").splitlines() if l.strip()]
    by_item = {}
    for r in rows:
        by_item.setdefault(r["_pair_id"], {})[r["condition"]] = r
    return by_item


def build_mapping(item_ids: list[str]) -> dict:
    """slot -> item_id -> {'X': 'direct'|'purpose_aware', 'Y': ...}, balanced A/B-in-X-position."""
    mapping = {}
    for slot in SLOTS:
        rng = random.Random(f"{SEED}|blind_eval|{slot}")
        mapping[slot] = {}
        for iid in item_ids:
            x_is_direct = rng.random() < 0.5
            mapping[slot][iid] = {
                "X": "direct" if x_is_direct else "purpose_aware",
                "Y": "purpose_aware" if x_is_direct else "direct",
            }
    return mapping


def sheet_guide(wb, slot):
    ws = wb.create_sheet("안내")
    lines = [
        "Pally 대화 목적 실험 — 최종 응답 블라인드 평가 (연구자 4인 독립, PM 설계서 §12-13)",
        "",
        f"평가자 슬롯: {slot}",
        "",
        "[ 중요 ]",
        "- 응답 X/Y는 어느 조건(A/B)인지 알 수 없게 익명화되어 있습니다. 유추하려 하지 마세요.",
        "- 다른 연구자와 채점 중 상의하지 마세요. 최초 독립 평가를 저장한 뒤에만 불일치를 검토합니다.",
        "- reference_response는 특정 문장 복제를 유도할 수 있어 이 화면에 제공하지 않습니다.",
        "",
        "[ 평가 순서 (설계서 §12.2) ]",
        "1. X, Y 각각 중대한 의미 오류(critical_error) 여부를 먼저 판정합니다.",
        "2. 의미·태도 보존(1-5), 목적 충족(1-5), 스타일 적절성(1-5), 대화 연속성(1-5), 관련성과 자연스러움(1-5)을 각각 매깁니다.",
        "3. '좋은 응답 통과'와 '승자'는 자동 계산됩니다(수식). 규칙과 본인 판단이 다르면 근거란에 이유를 적어주세요 — 기록되는 값은 규칙대로입니다.",
        "4. 근거(evidence)를 한 줄로 남깁니다.",
        "",
        "[ 중대한 오류(critical_error) 기준 (§13.1) — 하나라도 해당하면 Y ]",
        "- 문자적/비문자적 의미를 반대로 해석함",
        "- 칭찬·호의와 비난·적의를 반대로 해석함",
        "- 사용자의 핵심 요청과 반대되는 행동을 수행함",
        "- 입력에 없는 사실을 만들어 의미를 왜곡함",
        "- 안전·민감 상황에서 부적절한 농담/위험한 조언",
        "- 현재 문맥과 양립 불가능한 화제/전제를 응답의 핵심으로 삼음",
        "(단순한 문체 어색함, 짧은 응답 자체는 중대한 오류 아님)",
        "",
        "[ 좋은 응답 통과 기준 (§13.3, 자동 계산) ]",
        "critical_error=N 이고, 의미·태도보존>=4, 나머지 4개 항목 모두 >=3 이면 통과",
        "",
        "[ 승자 판정 규칙 (§13.4, 자동 계산) ]",
        "1. 한쪽만 통과 -> 통과한 쪽 승리",
        "2. 둘 다 미통과 -> both_bad",
        "3. 둘 다 통과, 5개 점수 합계 차이 >=3 -> 점수 높은 쪽 승리",
        "4. 둘 다 통과, 점수 합계 차이 <=2 -> tie",
    ]
    for i, line in enumerate(lines, start=1):
        ws.cell(row=i, column=1, value=line)
    ws["A1"].font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 100
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def sheet_scoring(wb, slot, items, responses, mapping):
    ws = wb.create_sheet("블라인드 채점")
    # NOTE: openpyxl's max_row does not count a row written via append([])
    # (no cells => nothing for max_row to see), even though the internal
    # append cursor still advances past it. Computing header_row from
    # max_row after a blank-row append undercounts by 1, which previously
    # caused the formula-writing loop below to also stomp on the header
    # row. Row numbers here are hardcoded to the fixed 3-row preamble
    # instead, so this can't drift silently again.
    ws.cell(row=1, column=1, value="reviewer_id ->")
    ws.cell(row=1, column=4, value="(여기에 본인 식별자 입력)")
    ws.cell(row=2, column=1, value="reviewer_slot ->")
    ws.cell(row=2, column=2, value=slot)

    base_cols = [
        "item_id", "context_before", "target_turn",
        "인간_합의_의미", "speaker_stance",
        "required_properties", "forbidden_errors",
        "continuation_anchors", "continuation_probe",
        "응답_X", "응답_Y",
    ]
    x_cols = ["X_critical_error"] + [f"X_{name}" for name, _ in DIMENSIONS] + ["X_좋은응답(자동)"]
    y_cols = ["Y_critical_error"] + [f"Y_{name}" for name, _ in DIMENSIONS] + ["Y_좋은응답(자동)"]
    tail_cols = ["승자(자동)", "근거", "annotator_id"]
    columns = base_cols + x_cols + y_cols + tail_cols

    header_row = 4
    for c, name in enumerate(columns, start=1):
        ws.cell(row=header_row, column=c, value=name)
    for cell in ws[header_row]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    col_index = {name: i + 1 for i, name in enumerate(columns)}

    def col_letter(name):
        return openpyxl.utils.get_column_letter(col_index[name])

    order = list(mapping[slot].keys())
    rng = random.Random(f"{SEED}|row_order|{slot}")
    rng.shuffle(order)

    first_data_row = header_row + 1
    for offset, iid in enumerate(order):
        r = first_data_row + offset
        it = items[iid]
        pair = responses[iid]
        cond_x = mapping[slot][iid]["X"]
        cond_y = mapping[slot][iid]["Y"]
        values = {
            "item_id": iid,
            "context_before": it.get("context_before") or "",
            "target_turn": it["target_turn"],
            "인간_합의_의미": it.get("intended_slang_sense") or "",
            "speaker_stance": it.get("speaker_stance") or "",
            "required_properties": it.get("required_properties") or "",
            "forbidden_errors": it.get("forbidden_errors") or "",
            "continuation_anchors": it.get("continuation_anchors") or "",
            "continuation_probe": it.get("continuation_probe") or "",
            "응답_X": pair[cond_x]["response"],
            "응답_Y": pair[cond_y]["response"],
        }
        for name, value in values.items():
            ws.cell(row=r, column=col_index[name], value=value)

    last_data_row = first_data_row + len(order) - 1

    # data validation: Y/N for critical_error, 1-5 for dimension scores
    dv_yn = DataValidation(type="list", formula1=f'"{",".join(YN)}"', allow_blank=True)
    dv_score = DataValidation(type="whole", operator="between", formula1=1, formula2=5, allow_blank=True)
    for side in ("X", "Y"):
        ce_col = col_letter(f"{side}_critical_error")
        dv_yn.add(f"{ce_col}{first_data_row}:{ce_col}{last_data_row}")
        for name, _ in DIMENSIONS:
            c = col_letter(f"{side}_{name}")
            dv_score.add(f"{c}{first_data_row}:{c}{last_data_row}")
    ws.add_data_validation(dv_yn)
    ws.add_data_validation(dv_score)

    # formulas: good_response (per side) and winner, row by row
    dim_letters = {side: [col_letter(f"{side}_{name}") for name, _ in DIMENSIONS] for side in ("X", "Y")}
    for r in range(first_data_row, last_data_row + 1):
        for side in ("X", "Y"):
            ce = f"{col_letter(f'{side}_critical_error')}{r}"
            meaning, purpose, style, cont, natural = (f"{c}{r}" for c in dim_letters[side])
            good_col = col_letter(f"{side}_좋은응답(자동)")
            formula = (
                f'=IF(OR({ce}="",{meaning}="",{purpose}="",{style}="",{cont}="",{natural}=""),"",'
                f'IF(AND({ce}="N",{meaning}>=4,{purpose}>=3,{style}>=3,{cont}>=3,{natural}>=3),"Y","N"))'
            )
            ws[f"{good_col}{r}"] = formula

        x_good = f"{col_letter('X_좋은응답(자동)')}{r}"
        y_good = f"{col_letter('Y_좋은응답(자동)')}{r}"
        x_sum = "+".join(f"{c}{r}" for c in dim_letters["X"])
        y_sum = "+".join(f"{c}{r}" for c in dim_letters["Y"])
        winner_col = col_letter("승자(자동)")
        ws[f"{winner_col}{r}"] = (
            f'=IF(OR({x_good}="",{y_good}=""),"",'
            f'IF(AND({x_good}="Y",{y_good}="N"),"X",'
            f'IF(AND({x_good}="N",{y_good}="Y"),"Y",'
            f'IF(AND({x_good}="N",{y_good}="N"),"both_bad",'
            f'IF(ABS(({x_sum})-({y_sum}))>=3,IF(({x_sum})>({y_sum}),"X","Y"),"tie")))))'
        )

    widths = {
        "item_id": 12, "context_before": 30, "target_turn": 30, "인간_합의_의미": 28,
        "speaker_stance": 12, "required_properties": 22, "forbidden_errors": 22,
        "continuation_anchors": 20, "continuation_probe": 24, "응답_X": 34, "응답_Y": 34,
        "근거": 30, "annotator_id": 14,
    }
    for name, w in widths.items():
        ws.column_dimensions[col_letter(name)].width = w
    for name in x_cols + y_cols:
        ws.column_dimensions[col_letter(name)].width = 10
    ws.column_dimensions[col_letter("승자(자동)")].width = 12

    for row in ws.iter_rows(min_row=first_data_row, max_row=last_data_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_workbook(slot, items, responses, mapping, out_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet_guide(wb, slot)
    sheet_scoring(wb, slot, items, responses, mapping)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    items = load_items()
    responses = load_responses()
    item_ids = sorted(responses.keys())
    print(f"loaded {len(item_ids)} final-test items with responses")

    mapping = build_mapping(item_ids)
    MAPPING_PATH.write_text(json.dumps(mapping, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote internal X/Y mapping -> {MAPPING_PATH} (DO NOT share with evaluators)")

    # sanity check: overall A-in-X balance
    x_is_direct_count = sum(
        1 for slot in SLOTS for iid in item_ids if mapping[slot][iid]["X"] == "direct"
    )
    total = len(SLOTS) * len(item_ids)
    print(f"A(direct) in X position: {x_is_direct_count}/{total} ({x_is_direct_count/total:.0%})")

    for slot in SLOTS:
        out_path = OUT_DIR / f"pally_purpose_v1_blind_eval_slot{slot}.xlsx"
        build_workbook(slot, items, responses, mapping, out_path)
        print(f"  wrote {out_path}")


if __name__ == "__main__":
    main()
