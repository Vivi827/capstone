# -*- coding: utf-8 -*-
"""Export the Pally purpose-aware 100-item labeling workbook (4 researcher slots).

Schema follows the PM final experiment design
(docs/AI_TACTIC/Pally_Communicative_Purpose_Aware_Experiment_Design_2026-09-22.md,
PR #87) Sections 6-9: 6 communicative-purpose categories, Formality/Humor/Curiosity
per rubric v2.2, plus the new fields Section 8.4 requires (intended_slang_sense,
literal_or_nonliteral, speaker_stance, reference_response/required_properties/
forbidden_errors, continuation_anchors/continuation_probe).

Item source: data/fixtures/pally_purpose_v1_items_100.jsonl (100 items: clear 30,
ambiguous-informal 40, surface-shortcut 30; train 80 / final_test 20 already
assigned and fixed -- labelers are NOT shown the split so it can't bias scoring).

All 4 researchers label all 100 items independently (per Section 9.1), each in a
different shuffled order.

Run from the repository root:
  python scripts/export_pally_purpose_v1_workbook.py
"""

from __future__ import annotations

import json
import random
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ITEMS_PATH = Path("data/fixtures/pally_purpose_v1_items_100.jsonl")
OUT_DIR = Path("data/fixtures")
SLOTS = ("A", "B", "C", "D")
SEED = 20260924

HEADER_FILL = PatternFill("solid", fgColor="DDDDDD")
BOLD = Font(bold=True)

PURPOSES = [
    ("answer_inform", "사실·뜻 답변 요구", "비교적 한정된 사실, 뜻 또는 정보를 요구하는가?"),
    ("explain_explore", "설명·탐색 요구", "이유, 원리, 비교, 해석 또는 가능성을 넓혀 묻는가?"),
    ("act_request", "행동·요청", "작성, 추천, 수정, 실행 또는 구체적 행동을 요구하는가?"),
    ("confirm_repair", "확인·명료화·수정", "이해 확인, 선택 확인, 오해 수정 또는 재설명을 요구하는가?"),
    ("acknowledge_empathize", "공유·평가·감정 반응", "경험, 상태, 감정 또는 평가에 대한 이해와 반응을 기대하는가?"),
    ("social_reciprocate", "사회적 화답·관계 유지", "인사, 감사, 칭찬, 감탄 등 관계적 화답과 대화 지속이 중심인가?"),
]
PURPOSE_CODES = [p[0] for p in PURPOSES]

LANGUAGE_FORMS = [
    ("standard_or_unmarked", "뚜렷한 오류·비표준성이 없음"),
    ("likely_learner_error", "의미는 전달되나 비의도적 학습자 오류로 보임"),
    ("conventional_nonstandard", "ain't, 축약, 구어 문법처럼 관습화된 비표준형"),
    ("playful_deviation", "의도적 철자·문법 파괴, 밈 문형, 과장된 반복 등 유희적 변형의 근거가 있음"),
    ("intent_unclear", "오류인지 의도적 변형인지 단일 발화로 판정 불가"),
]
LANGUAGE_FORM_CODES = [f[0] for f in LANGUAGE_FORMS]

LITERAL_NONLITERAL = ["literal", "nonliteral", "mixed_or_ambiguous"]
STANCE = ["긍정", "부정", "중립", "혼합"]
CASE_TYPES = ["clear", "boundary", "context_insufficient"]
SCORE_STATUS = ["final", "provisional"]
YN = ["Y", "N"]

AXES = ("Formality", "Humor", "Curiosity")

AXIS_DEFINITIONS = {
    "Formality": ("현재 발화의 구어적 존중, 체면 고려, 직접성 완화 및 register",
                  "`please`, 문장 길이, 문법 정확성만으로 상승시키지 않음"),
    "Humor": ("현재 발화에서 지각되는 유희적 프레이밍, 비문자성, 불일치 및 검수된 밈·슬랭의 유희적 사용",
              "슬랭 출현이나 문법 오류만으로 상승시키지 않음"),
    "Curiosity": ("현재 발화가 정보, 설명, 관점 또는 가능성을 얼마나 명시적으로 요구하는지",
                  "물음표나 질문형만으로 높은 점수를 부여하지 않음"),
}

GOOD_RESPONSE_CRITERIA = [
    "슬랭의 문자적·비문자적 의미와 화자의 태도를 정확히 보존한다.",
    "현재 발화가 요구하는 핵심 반응을 수행한다.",
    "사용자 스타일과 조화를 이루되 공격성, 과도한 친밀감, 부적절한 농담을 그대로 모방하지 않는다.",
    "동일한 의미 맥락에서 이어질 수 있는 구체적인 반응 또는 후속 화제를 제공한다.",
    "간결하고 자연스러우며 새로운 사실을 근거 없이 만들어내지 않는다.",
]


def load_items() -> list[dict]:
    items = []
    with open(ITEMS_PATH, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line:
                items.append(json.loads(line))
    return items


def sheet_guide(wb: openpyxl.Workbook, slot: str, n_items: int, seed: int) -> None:
    ws = wb.create_sheet("안내")
    lines = [
        "Pally 대화 목적 선행형 실험 — 인간 라벨링 안내 (100개, 연구자 4인 독립)",
        "",
        f"평가자 슬롯: {slot}   /   문항 수: {n_items}   /   순서 seed: {seed}",
        "",
        "[ 목적 ]",
        "조건 A(Direct)와 조건 B(Purpose-aware)를 비교하는 실험의 인간 정답 라벨을 만든다.",
        "여기서 매기는 값은 두 조건 모두를 비교하는 공통 정답이며, 이 라벨링 자체는 세 번째 모델 조건이 아니다.",
        "Energy와 Intimacy는 이번 실험에서 측정하지 않는다 (실제 음성·사용 이력이 없음).",
        "",
        "[ 채점 순서 (설계서 §9.1, §12.2 기준) ]",
        "1. '대화 목적' 시트에서 6개 목적 범주를 먼저 읽는다.",
        "2. '축 정의' 시트로 Formality/Humor/Curiosity 판단 기준과 금지되는 shortcut을 확인한다.",
        "3. '채점 시트'에서 각 항목(context_before + target_turn)에 대해:",
        "   a. 주 대화 목적 1개 선택 (필요하면 보조 목적 1개, 모호하면 purpose_ambiguous=Y)",
        "   b. language_form, literal_or_nonliteral, speaker_stance 선택",
        "   c. intended_slang_sense: 의도된 문자적/비문자적 의미를 짧게 서술",
        "   d. Formality/Humor/Curiosity 점수(0~100)와 case_type, confidence(0~4), 근거",
        "   e. reference_response(예시 응답 하나), required_properties(좋은 응답의 필수 속성),",
        "      forbidden_errors(발생하면 안 되는 오류) — 표현이 달라도 속성을 충족하면 정답으로 인정",
        "   f. continuation_anchors(허용 가능한 후속 화제/반응), continuation_probe(평가용 다음 사용자 턴 — 모델에는 안 보여줌)",
        "4. 상단 reviewer_id 칸에 본인 식별자를 입력한다.",
        "",
        "[ 규칙 ]",
        "- 이 항목이 train인지 final_test인지 여기서는 알 수 없다 (일부러 안 보여줌 — 모델·평가 기준에 영향 주지 않기 위함).",
        "- 다른 평가자의 라벨, 출처(어느 코퍼스인지), 모델 예측값을 보지 않는다. 현재 항목(context_before+target_turn)만 본다.",
        "- 0은 축의 낮은 끝을 관찰했다는 뜻이고, null은 해당 입력으로 측정하지 않았다는 뜻이다 (Energy/Intimacy만 null).",
        "- 문법 오류 자체는 세 축 모두 직접 점수에 반영하지 않는다 (language_form으로 별도 기록).",
        "- joke/humorous/playful을 목적 범주로 쓰지 않는다 (Humor 정답이 목적 선택에 순환 반영되는 것을 막기 위함).",
        "- 여러 해석이 가능하면 case_type을 boundary 또는 context_insufficient로 표시하고, 그래도 가장 지지되는 잠정 점수를 기록한다 (판단 불가 시 50점 기본값 금지).",
        "- 다른 평가자와 상의하지 말고 완전히 독립적으로 채점한다.",
        "",
        "[ 참고: '좋은 응답'의 정의 (설계서 §7.2, required_properties 작성에 참고) ]",
    ] + [f"  - {c}" for c in GOOD_RESPONSE_CRITERIA]
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 110
    for cell in ws["A"]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(bold=True, size=13)


def sheet_purposes(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("대화 목적·문법기준")
    ws.append(["대화 목적 6종 (설계서 §6)"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["코드", "한국어명", "판정 질문"])
    for cell in ws[2]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    for code, label, question in PURPOSES:
        ws.append([code, label, question])
    ws.append([])
    ws.append(["주의: joke/humorous/playful은 목적 범주로 만들지 않는다 (Humor 순환 판단 방지)."])
    row0 = ws.max_row + 2
    ws.cell(row=row0, column=1, value="문법 오류·비표준 표현 (language_form)").font = Font(bold=True, size=12)
    ws.append(["코드", "정의"])
    for cell in ws[row0 + 1]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    for code, definition in LANGUAGE_FORMS:
        ws.append([code, definition])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 50
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def sheet_axes(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("축 정의")
    ws.append(["Formality / Humor / Curiosity 정의 (설계서 §7.1)"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(["축", "본 실험의 정의", "금지되는 shortcut"])
    for cell in ws[2]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    for axis in AXES:
        definition, shortcut = AXIS_DEFINITIONS[axis]
        ws.append([axis, definition, shortcut])
    ws.append([])
    ws.append(["세 축은 각각 0-100점, v2.2 구성요소 가중치와 5점 단위 반올림 유지. 세 축을 하나의 합산 점수로 만들지 않는다."])
    ws.append(["Energy/Intimacy는 0이 아니라 null."])
    r = ws.max_row + 2
    ws.cell(row=r, column=1, value="경계 사례·점수 상태").font = Font(bold=True, size=12)
    r += 1
    ws.append(["case_type", "의미"])
    for cell in ws[r + 1]:
        cell.font = BOLD
        cell.fill = HEADER_FILL
    ws.append(["clear", "현재 항목에서 우세한 의미와 축 점수를 안정적으로 판단 가능"])
    ws.append(["boundary", "복수 해석이 가능하지만 우세한 판단은 가능"])
    ws.append(["context_insufficient", "문맥 없이는 우세한 의미나 의도를 결정하기 어려움"])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 55
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


BASE_COLS = [
    "item_id", "context_before", "target_turn",
    "purpose_primary", "purpose_secondary", "purpose_ambiguous",
    "language_form", "literal_or_nonliteral", "speaker_stance", "intended_slang_sense",
]
AXIS_COL_GROUPS = {
    "Formality": ["Formality", "Formality_case_type", "Formality_confidence", "Formality_rationale"],
    "Humor": ["Humor", "Humor_case_type", "Humor_confidence", "Humor_rationale"],
    "Curiosity": ["Curiosity", "Curiosity_case_type", "Curiosity_confidence", "Curiosity_rationale"],
}
TAIL_COLS = [
    "Energy", "Intimacy",
    "reference_response", "required_properties", "forbidden_errors",
    "continuation_anchors", "continuation_probe",
    "annotator_id", "메모(선택)",
]


def sheet_scoring(wb: openpyxl.Workbook, slot: str, items: list[dict], seed: int) -> None:
    ws = wb.create_sheet("채점 시트")
    ws.append(["reviewer_id →", "", "", "(여기에 본인 식별자 입력)"])
    ws.append(["reviewer_slot →", slot])
    ws.append([])

    columns = (
        BASE_COLS
        + AXIS_COL_GROUPS["Formality"]
        + AXIS_COL_GROUPS["Humor"]
        + AXIS_COL_GROUPS["Curiosity"]
        + TAIL_COLS
    )
    header_row = ws.max_row + 1
    ws.append(columns)
    for cell in ws[header_row]:
        cell.font = BOLD
        cell.fill = HEADER_FILL

    order = list(range(len(items)))
    random.Random(f"{seed}|pally_purpose_v1|{slot}").shuffle(order)

    col_index = {name: i + 1 for i, name in enumerate(columns)}

    for idx in order:
        item = items[idx]
        row = [""] * len(columns)
        row[col_index["item_id"] - 1] = item["item_id"]
        row[col_index["context_before"] - 1] = item.get("context_before") or ""
        row[col_index["target_turn"] - 1] = item["target_turn"]
        row[col_index["Energy"] - 1] = "null"
        row[col_index["Intimacy"] - 1] = "null"
        ws.append(row)

    first_data_row = header_row + 1
    last_data_row = header_row + len(items)

    def col_letter(name: str) -> str:
        return openpyxl.utils.get_column_letter(col_index[name])

    def col_range(name: str) -> str:
        letter = col_letter(name)
        return f"{letter}{first_data_row}:{letter}{last_data_row}"

    dv_purpose = DataValidation(type="list", formula1=f'"{",".join(PURPOSE_CODES)}"', allow_blank=True)
    dv_purpose2 = DataValidation(type="list", formula1=f'"{",".join(PURPOSE_CODES)}"', allow_blank=True)
    dv_yn = DataValidation(type="list", formula1=f'"{",".join(YN)}"', allow_blank=True)
    dv_langform = DataValidation(type="list", formula1=f'"{",".join(LANGUAGE_FORM_CODES)}"', allow_blank=True)
    dv_literal = DataValidation(type="list", formula1=f'"{",".join(LITERAL_NONLITERAL)}"', allow_blank=True)
    dv_stance = DataValidation(type="list", formula1=f'"{",".join(STANCE)}"', allow_blank=True)
    dv_case_type = DataValidation(type="list", formula1=f'"{",".join(CASE_TYPES)}"', allow_blank=True)
    dv_confidence = DataValidation(type="whole", operator="between", formula1=0, formula2=4, allow_blank=True)
    dv_score = DataValidation(type="whole", operator="between", formula1=0, formula2=100, allow_blank=True)

    dv_purpose.add(col_range("purpose_primary"))
    dv_purpose2.add(col_range("purpose_secondary"))
    dv_yn.add(col_range("purpose_ambiguous"))
    dv_langform.add(col_range("language_form"))
    dv_literal.add(col_range("literal_or_nonliteral"))
    dv_stance.add(col_range("speaker_stance"))

    for axis in AXES:
        dv_score.add(col_range(axis))
        dv_case_type.add(col_range(axis + "_case_type"))
        dv_confidence.add(col_range(axis + "_confidence"))

    for dv in (dv_purpose, dv_purpose2, dv_yn, dv_langform, dv_literal, dv_stance, dv_case_type, dv_confidence, dv_score):
        ws.add_data_validation(dv)

    widths = {
        "item_id": 16, "context_before": 34, "target_turn": 46,
        "purpose_primary": 20, "purpose_secondary": 20, "purpose_ambiguous": 10,
        "language_form": 22, "literal_or_nonliteral": 16, "speaker_stance": 10,
        "intended_slang_sense": 34,
    }
    for axis in AXES:
        widths.update({
            axis: 8, f"{axis}_case_type": 18, f"{axis}_confidence": 12, f"{axis}_rationale": 30,
        })
    widths.update({
        "Energy": 8, "Intimacy": 8,
        "reference_response": 34, "required_properties": 30, "forbidden_errors": 30,
        "continuation_anchors": 30, "continuation_probe": 30,
        "annotator_id": 14, "메모(선택)": 26,
    })
    for name, w in widths.items():
        ws.column_dimensions[col_letter(name)].width = w

    for row in ws.iter_rows(min_row=first_data_row, max_row=last_data_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_workbook(slot: str, items: list[dict], seed: int, out_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheet_guide(wb, slot, len(items), seed)
    sheet_purposes(wb)
    sheet_axes(wb)
    sheet_scoring(wb, slot, items, seed)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main() -> None:
    items = load_items()
    print(f"loaded {len(items)} items")
    for slot in SLOTS:
        out_path = OUT_DIR / f"pally_purpose_v1_scoring_slot{slot}.xlsx"
        build_workbook(slot, items, SEED, out_path)
        print(f"  wrote {out_path}")


if __name__ == "__main__":
    main()
